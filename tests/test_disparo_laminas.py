"""Testes do Disparo de Lâminas (núcleo, sem Outlook nem tela)."""

from __future__ import annotations

import queue
import sys
from datetime import datetime
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "disparo_laminas"))

from laminas import config as cfgmod  # noqa: E402
from laminas.base import (DUPLICADO, INVALIDO, JA_ENVIADO, PENDENTE, Contato,  # noqa: E402
                          analisar, limpar_texto, registros)
from laminas.campanhas import (CAMPANHA_GERAL, lamina_do_envio, listar_campanhas,  # noqa: E402
                               listar_laminas)
from laminas.disparo import Disparo  # noqa: E402
from laminas.envio import EnviadorSimulacao  # noqa: E402
from laminas.historico import ENVIADO, ERRO, Historico  # noqa: E402
from laminas.mensagem import (corpo_html, dimensoes_imagem, inserir_no_inicio_do_body,  # noqa: E402
                              montar, preencher)

PNG_1x1 = bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000320000002580806000000"
    "9a768270000000017352474200aece1ce90000000d49444154789c636000"
    "0000000200015e0f2d1d0000000049454e44ae426082"
)  # cabeçalho PNG 800x600


def _planilha(tmp_path, linhas, aba="A ENVIAR"):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = aba
    for linha in linhas:
        ws.append(linha)
    caminho = tmp_path / "base.xlsx"
    wb.save(caminho)
    return str(caminho)


@pytest.fixture()
def hist(tmp_path):
    return Historico(tmp_path / "h.db")


# -- base / duplicados ---------------------------------------------------------

def test_limpar_texto_remove_lixo_de_exportacao():
    assert limpar_texto("_x000D_\n\tBDO Brasil - BDO_x000D_\n") == "BDO Brasil - BDO"
    assert limpar_texto("\n            Dantas Distribuidora\n        ") == "Dantas Distribuidora"


def test_registros_reconhece_titulos_da_macro(tmp_path):
    caminho = _planilha(tmp_path, [
        ["Empresa", "Nome Completo", "E-mail", None],
        ["  Acme  ", "Ana Souza", " Ana@Acme.com.br ", "=VLOOKUP(...)"],
    ])
    regs = registros(caminho, "A ENVIAR")
    assert regs == [{"linha": 2, "empresa": "Acme", "nome": "Ana Souza",
                     "email": "Ana@Acme.com.br", "copia": "", "data": None, "status": ""}]


def test_colunas_em_qualquer_ordem_e_com_outros_titulos(tmp_path):
    caminho = _planilha(tmp_path, [
        ["Lista de prospecção - outubro"],  # linha de título acima do cabeçalho
        [],
        ["Cargo", "E-mail Principal", "Mensagem (parte de cima do e-mail)", "Razão Social", "Contato"],
        ["Diretora", "ana@acme.com", "texto", "Acme Ltda", "Ana Souza"],
    ])
    from laminas.base import ler_base
    regs, cab = ler_base(caminho, "A ENVIAR")
    assert (regs[0]["nome"], regs[0]["empresa"], regs[0]["email"]) == ("Ana Souza", "Acme Ltda", "ana@acme.com")
    assert cab.linha_titulo == 3
    assert cab.descricao("nome") == "E - Contato" and cab.descricao("email") == "B - E-mail Principal"


def test_primeiro_nome_mais_sobrenome_e_nome_da_empresa(tmp_path):
    caminho = _planilha(tmp_path, [
        ["Sobrenome", "Nome da Empresa", "Primeiro Nome", "Email"],
        ["Souza", "Acme", "Ana", "ana@acme.com"],
    ])
    reg = registros(caminho)[0]
    assert (reg["nome"], reg["empresa"]) == ("Ana Souza", "Acme")


def test_planilha_sem_titulos_e_escolha_manual_das_colunas(tmp_path):
    caminho = _planilha(tmp_path, [
        ["Acme", "Ana Souza", "ana@acme.com"],
        ["Beta", "Bia Lima", "bia@beta.com"],
    ])
    from laminas.base import ler_base
    regs, cab = ler_base(caminho)
    assert cab.linha_titulo == 0 and regs[0]["email"] == "ana@acme.com" and regs[0]["nome"] == ""
    regs, cab = ler_base(caminho, colunas={"empresa": 0, "nome": 1})
    assert len(regs) == 2 and (regs[1]["nome"], regs[1]["empresa"]) == ("Bia Lima", "Beta")


def test_titulos_genericos_nao_viram_contato(tmp_path):
    caminho = _planilha(tmp_path, [["Col1", "Col2", "Col3"], ["Gama", "Carla", "carla@gama.com"]])
    from laminas.base import ler_base
    regs, cab = ler_base(caminho, colunas={"empresa": 0, "nome": 1})
    assert cab.linha_titulo == 1 and cab.descricao("nome") == "B - Col2"
    assert [(r["linha"], r["nome"], r["empresa"]) for r in regs] == [(2, "Carla", "Gama")]


def test_registros_csv_com_ponto_e_virgula(tmp_path):
    caminho = tmp_path / "base.csv"
    caminho.write_text("Empresa;Nome;E-mail;Cópia\nAcme;Ana;ana@acme.com;S\n", encoding="utf-8")
    regs = registros(str(caminho))
    assert regs[0]["email"] == "ana@acme.com" and regs[0]["copia"] == "S"


def test_analisar_aponta_duplicados_ja_enviados_e_invalidos():
    regs = [
        {"linha": 2, "empresa": "A", "nome": "Ana", "email": "ana@a.com", "copia": "S"},
        {"linha": 3, "empresa": "A", "nome": "Ana 2", "email": "ANA@a.com ", "copia": ""},
        {"linha": 4, "empresa": "B", "nome": "Bia", "email": "bia@b.com", "copia": ""},
        {"linha": 5, "empresa": "C", "nome": "Caio", "email": "caio-sem-arroba", "copia": ""},
        {"linha": 6, "empresa": "", "nome": "Duda", "email": "duda@d.com", "copia": ""},
        {"linha": 7, "empresa": "A", "nome": "Ana 3", "email": "ana@a.com", "copia": ""},
    ]
    contatos, resumo = analisar(regs, {"bia@b.com": "2026-04-06 13:47:34"})
    sit = {c.linha: c.situacao for c in contatos}
    assert sit == {2: PENDENTE, 3: DUPLICADO, 4: JA_ENVIADO, 5: INVALIDO, 6: INVALIDO, 7: DUPLICADO}
    assert resumo.duplicados == {"ana@a.com": [2, 3, 7]}
    assert contatos[0].enviar and contatos[0].cc  # coluna Cópia = S já marca CC
    assert not contatos[2].enviar and "06/04/2026" in contatos[2].detalhe
    assert "empresa" in contatos[4].detalhe


# -- histórico -----------------------------------------------------------------

def test_historico_registra_conta_dia_e_dedup(hist):
    hist.registrar(email="Ana@A.com", nome="Ana", status=ENVIADO, data_hora="2026-10-27 09:00:00")
    hist.registrar(email="bia@b.com", status=ERRO, data_hora="2026-10-27 09:03:00")
    assert hist.enviados_no_dia(datetime(2026, 10, 27).date()) == 1
    assert hist.data_envio("ana@a.com") == "2026-10-27 09:00:00"
    assert hist.data_envio("bia@b.com") is None  # erro não conta como enviado
    assert set(hist.ja_enviados()) == {"ana@a.com"}


def test_historico_importa_relatorio_antigo_e_exporta(tmp_path, hist):
    caminho = _planilha(tmp_path, [
        ["Empresa", "Nome do Cliente", "E-mail Enviado", "Data", "Hora", "Status"],
        ["X", "Ana", "ana@a.com", datetime(2026, 4, 6), None, "Enviado"],
        ["X", "Bia", "bia@b.com", datetime(2026, 4, 6), None, "Erro"],
        ["X", "Ana", "ana@a.com", datetime(2026, 4, 7), None, "Enviado"],
    ], aba="RELATORIO")
    novos, ignorados = hist.importar(registros(caminho, "RELATORIO"), origem="teste")
    assert (novos, ignorados) == (1, 2)
    assert hist.data_envio("ana@a.com") == "2026-04-06 00:00:00"
    assert hist.enviados_no_dia(datetime(2026, 4, 6).date()) == 0  # importado não gasta limite

    destino = tmp_path / "enviados.xlsx"
    assert hist.exportar_xlsx(destino) == 1
    from openpyxl import load_workbook
    ws = load_workbook(destino)["ENVIADOS"]
    assert ws["D2"].value == "ana@a.com" and ws["J2"].value == "Importado"


# -- mensagem ------------------------------------------------------------------

def test_preencher_campos_novos_e_da_macro():
    c = Contato(linha=2, empresa="Acme", nome="Ana Souza", email="ana@acme.com")
    manha, tarde = datetime(2026, 1, 1, 9), datetime(2026, 1, 1, 15)
    assert preencher("Olá {nome}, {saudacao}! ({empresa})", c, manha) == "Olá Ana Souza, bom dia! (Acme)"
    assert preencher("Olá <<Nome do Cliente>>, <diatarde>! {primeiro_nome}", c, tarde) == \
        "Olá Ana Souza, boa tarde! Ana"
    assert preencher("{desconhecido}", c, manha) == "{desconhecido}"


def test_montar_so_poe_cc_nos_marcados(tmp_path):
    msg = cfgmod.mensagem_da_campanha(cfgmod.DEFAULTS, "corporativo")
    com = Contato(2, "Acme", "Ana", "ana@acme.com", cc=True)
    sem = Contato(3, "Acme", "Bia", "bia@acme.com", cc=False)
    assert montar(com, msg, imagem=None, cc="ricardo@x.com").cc == "ricardo@x.com"
    assert montar(sem, msg, imagem=None, cc="ricardo@x.com").cc == ""
    assert montar(com, msg, imagem=None).assunto.startswith("CONFIRMANDO REUNIÃO com Ana da Acme")


def test_corpo_html_limita_largura_e_assinatura(tmp_path):
    img = tmp_path / "l.png"
    img.write_bytes(PNG_1x1)
    assert dimensoes_imagem(img) == (800, 600)
    c = Contato(2, "Acme", "Ana <b>", "ana@acme.com")
    email = montar(c, {"assunto": "x", "texto_inicio": "Olá {nome}\nlinha 2", "texto_fim": "Tchau"},
                   imagem=img)
    html = corpo_html(email, "cid:abc", 560)
    assert 'width="560" height="420"' in html and "Ana &lt;b&gt;" in html and "<br>" in html
    assinado = inserir_no_inicio_do_body('<html><body lang=PT-BR><p>Assinatura</p></body></html>', "<p>Oi</p>")
    assert assinado.index("Oi") < assinado.index("Assinatura")


# -- campanhas -------------------------------------------------------------------

def test_campanhas_e_rodizio_de_laminas(tmp_path):
    (tmp_path / "CORPORATIVO").mkdir()
    (tmp_path / "EVENTOS").mkdir()
    (tmp_path / "VAZIA").mkdir()
    for nome in ("a.jpeg", "b.png"):
        (tmp_path / "CORPORATIVO" / nome).write_bytes(PNG_1x1)
    (tmp_path / "EVENTOS" / "e.jpg").write_bytes(PNG_1x1)
    (tmp_path / "solta.png").write_bytes(PNG_1x1)
    assert listar_campanhas(str(tmp_path)) == ["CORPORATIVO", "EVENTOS", CAMPANHA_GERAL]
    laminas = listar_laminas(str(tmp_path), "CORPORATIVO")
    assert {p.name for p in laminas} == {"a.jpeg", "b.png"}
    assert [lamina_do_envio(laminas, i, True) for i in range(3)] == [laminas[0], laminas[1], laminas[0]]
    assert lamina_do_envio(laminas, 5, False) == laminas[0]
    assert listar_campanhas(str(tmp_path / "nao-existe")) == []


def test_config_mensagem_por_campanha(tmp_path):
    caminho = tmp_path / "config.json"
    cfg = cfgmod.carregar(caminho)
    cfg["mensagens"]["EVENTOS"] = {"assunto": "Eventos", "texto_inicio": "", "texto_fim": ""}
    cfg["limite_diario"] = 50
    cfgmod.salvar(cfg, caminho)
    lido = cfgmod.carregar(caminho)
    assert lido["limite_diario"] == 50 and lido["intervalo_minutos"] == 3.0
    assert cfgmod.mensagem_da_campanha(lido, "EVENTOS")["assunto"] == "Eventos"
    assert cfgmod.mensagem_da_campanha(lido, "CORPORATIVO")["assunto"].startswith("CONFIRMANDO")


# -- disparo ---------------------------------------------------------------------

class _Relogio:
    """Relógio falso: ``espera`` avança o tempo sem dormir."""

    def __init__(self):
        self.t = 0.0
        self.esperas: list[float] = []

    def __call__(self):
        return self.t


class _Enviador:
    def __init__(self, falhar=()):
        self.enviados = []
        self.falhar = set(falhar)

    def enviar(self, email):
        if email.para in self.falhar:
            raise RuntimeError("caixa cheia")
        self.enviados.append(email)


def _disparo(contatos, hist, enviador, relogio, **kw):
    eventos = queue.Queue()
    params = dict(historico=hist, criar_enviador=lambda: enviador,
                  mensagem={"assunto": "Oi {nome}", "texto_inicio": "", "texto_fim": ""},
                  campanha="CORPORATIVO", laminas=[Path("a.png"), Path("b.png")],
                  alternar_laminas=True, email_copia="ricardo@x.com", copia_oculta="",
                  limite_diario=80, intervalo_seg=180, eventos=eventos,
                  agora=lambda: datetime(2026, 10, 27, 10, 0), relogio=relogio)
    params.update(kw)
    d = Disparo(contatos, **params)

    def esperar(segundos):  # substitui o Event.wait real
        relogio.esperas.append(segundos)
        relogio.t += segundos
        return False
    d._parar.wait = esperar
    return d, eventos


def _eventos(fila):
    saida = []
    while not fila.empty():
        saida.append(fila.get())
    return saida


def test_disparo_intervalo_cc_rodizio_e_registro(hist):
    contatos = [Contato(i, "Acme", f"N{i}", f"n{i}@acme.com", enviar=True, cc=(i < 2))
                for i in range(4)]
    relogio = _Relogio()
    enviador = _Enviador()
    d, fila = _disparo(contatos, hist, enviador, relogio)
    d.run()
    assert [e.para for e in enviador.enviados] == [c.email for c in contatos]
    assert [e.cc for e in enviador.enviados] == ["ricardo@x.com", "ricardo@x.com", "", ""]
    assert [e.imagem.name for e in enviador.enviados] == ["a.png", "b.png", "a.png", "b.png"]
    assert relogio.esperas == [180, 180, 180]  # 1º sai na hora, depois 3 min entre cada
    assert hist.enviados_no_dia(datetime(2026, 10, 27).date()) == 4
    tipo, dados = _eventos(fila)[-1]
    assert tipo == "fim" and dados["enviados"] == 4


def test_disparo_respeita_limite_diario_contando_envios_anteriores(hist):
    # O histórico grava a hora real do envio, então o "dia" do teste é hoje.
    hoje = datetime.now().strftime("%Y-%m-%d")
    for i in range(78):
        hist.registrar(email=f"antes{i}@x.com", status=ENVIADO, data_hora=f"{hoje} 00:00:{i % 60:02d}")
    contatos = [Contato(i, "Acme", f"N{i}", f"n{i}@acme.com", enviar=True) for i in range(5)]
    enviador = _Enviador()
    d, fila = _disparo(contatos, hist, enviador, _Relogio(), agora=datetime.now)
    d.run()
    assert len(enviador.enviados) == 2
    tipo, dados = _eventos(fila)[-1]
    assert "Limite diário de 80" in dados["motivo"]


def test_disparo_pula_ja_enviado_salvo_reenvio_autorizado_e_registra_erro(hist):
    hist.registrar(email="velho@x.com", status=ENVIADO, data_hora="2026-01-01 10:00:00")
    hist.registrar(email="forca@x.com", status=ENVIADO, data_hora="2026-01-01 10:00:00")
    contatos = [
        Contato(2, "A", "Velho", "velho@x.com", enviar=True),
        Contato(3, "A", "Força", "forca@x.com", enviar=True, forcar=True),
        Contato(4, "A", "Falha", "falha@x.com", enviar=True),
    ]
    enviador = _Enviador(falhar={"falha@x.com"})
    d, fila = _disparo(contatos, hist, enviador, _Relogio())
    d.run()
    assert [e.para for e in enviador.enviados] == ["forca@x.com"]
    tipos = [t for t, _ in _eventos(fila)]
    assert "pulado" in tipos and "erro" in tipos
    erros = [r for r in hist.listar() if r["status"] == ERRO]
    assert erros[0]["email"] == "falha@x.com" and erros[0]["erro"] == "caixa cheia"


def test_disparo_parar_interrompe_na_espera(hist):
    contatos = [Contato(i, "A", f"N{i}", f"n{i}@a.com", enviar=True) for i in range(3)]
    enviador = _Enviador()
    d, fila = _disparo(contatos, hist, enviador, _Relogio())
    d._parar.wait = lambda s: True  # usuário clicou em Parar durante a espera
    d.run()
    assert len(enviador.enviados) == 1
    assert "interrompido" in _eventos(fila)[-1][1]["motivo"]


def test_enviador_simulacao_gera_html(tmp_path):
    img = tmp_path / "l.png"
    img.write_bytes(PNG_1x1)
    c = Contato(2, "Acme", "Ana", "ana@acme.com", cc=True)
    email = montar(c, {"assunto": "Oi", "texto_inicio": "Olá", "texto_fim": "Até"}, imagem=img, cc="r@x.com")
    EnviadorSimulacao(tmp_path / "sim").enviar(email)
    arquivos = list((tmp_path / "sim").glob("*.html"))
    assert len(arquivos) == 1
    conteudo = arquivos[0].read_text(encoding="utf-8")
    assert "CC: r@x.com" in conteudo and 'src="l.png"' in conteudo


# -- apresentação anexada por campanha ----------------------------------------

def test_apresentacao_por_campanha_e_pasta_nao_vira_campanha(tmp_path):
    from laminas.campanhas import listar_anexos, pasta_apresentacoes, sugerir_anexo
    for camp in ("CORPORATIVO", "EVENTOS"):
        (tmp_path / camp).mkdir()
        (tmp_path / camp / "l.png").write_bytes(PNG_1x1)
    apres = tmp_path / "APRESENTAÇÃO"
    apres.mkdir()
    (apres / "capa.png").write_bytes(PNG_1x1)  # imagem lá dentro não cria campanha
    (apres / "Apresentação Institucional Tivolitur.pdf").write_bytes(b"%PDF")
    (apres / "Apresentação Eventos Tivolitur.pptx").write_bytes(b"PK")
    (apres / "~$Apresentação Eventos Tivolitur.pptx").write_bytes(b"")
    assert listar_campanhas(str(tmp_path)) == ["CORPORATIVO", "EVENTOS"]
    pasta = pasta_apresentacoes(str(tmp_path))
    assert pasta == apres
    anexos = listar_anexos(pasta)
    assert [p.name for p in anexos] == ["Apresentação Eventos Tivolitur.pptx",
                                        "Apresentação Institucional Tivolitur.pdf"]
    assert sugerir_anexo(anexos, "CORPORATIVO").name == "Apresentação Institucional Tivolitur.pdf"
    assert sugerir_anexo(anexos, "Eventos").name == "Apresentação Eventos Tivolitur.pptx"
    assert sugerir_anexo(anexos, "INCENTIVO") is None


def test_disparo_anexa_apresentacao_e_registra(hist, tmp_path):
    pdf = tmp_path / "Institucional.pdf"
    pdf.write_bytes(b"%PDF")
    contatos = [Contato(2, "Acme", "Ana", "ana@acme.com", enviar=True)]
    enviador = _Enviador()
    d, _ = _disparo(contatos, hist, enviador, _Relogio(), anexos=[pdf])
    d.run()
    assert enviador.enviados[0].anexos == [pdf]
    assert hist.listar()[0]["anexo"] == "Institucional.pdf"

    EnviadorSimulacao(tmp_path / "sim").enviar(enviador.enviados[0])
    html = next((tmp_path / "sim").glob("*.html")).read_text(encoding="utf-8")
    assert "Anexos: Institucional.pdf" in html


def test_historico_antigo_ganha_coluna_anexo(tmp_path):
    import sqlite3
    caminho = tmp_path / "antigo.db"
    with sqlite3.connect(caminho) as con:
        con.execute("CREATE TABLE envios (id INTEGER PRIMARY KEY AUTOINCREMENT, data_hora TEXT,"
                    " empresa TEXT, nome TEXT, email TEXT NOT NULL, email_norm TEXT NOT NULL, cc TEXT,"
                    " campanha TEXT, lamina TEXT, assunto TEXT, status TEXT NOT NULL, erro TEXT)")
        con.execute("INSERT INTO envios (email, email_norm, status) VALUES ('a@a.com', 'a@a.com', 'Enviado')")
    h = Historico(caminho)
    h.registrar(email="b@b.com", anexo="Eventos.pptx")
    assert {r["email"]: r["anexo"] for r in h.listar()} == {"a@a.com": None, "b@b.com": "Eventos.pptx"}


# -- erros do Outlook ------------------------------------------------------------

class _ComErro(Exception):
    """Imita ``pywintypes.com_error`` (hresult, texto, excepinfo, argerr)."""


def test_descrever_erro_com_traduz_codigo_e_etapa():
    from laminas.envio import ErroEnvio, _etapa, descrever_erro
    erro = _ComErro(-2147024809, "Parâmetro incorreto.", None, None)
    assert descrever_erro(erro) == ("Parâmetro incorreto. [0x80070057] — o Outlook recusou "
                                    "um dos dados (parâmetro incorreto)")
    com_info = _ComErro(-2147352567, "Exceção.", (4096, "Microsoft Outlook",
                                                  "Não é possível localizar este arquivo.",
                                                  None, 0, -2147024894), None)
    assert descrever_erro(com_info).startswith("Não é possível localizar este arquivo. [0x80070002]")

    def falha():
        raise erro
    with pytest.raises(ErroEnvio, match=r"^Falha ao anexar a apresentação \(x.pdf\): Parâmetro incorreto"):
        _etapa("anexar a apresentação (x.pdf)", falha)


def test_anexo_recusado_tenta_copia_local(tmp_path, monkeypatch):
    import tempfile
    from laminas.envio import EnviadorOutlook
    monkeypatch.setattr(tempfile, "tempdir", str(tmp_path / "tmp"))
    pdf = tmp_path / "TIVOLITUR - APRESENTAÇÃO INSTITUCIONAL.pdf"
    pdf.write_bytes(b"%PDF")
    tentativas = []

    class Anexos:
        def Add(self, caminho):
            tentativas.append(caminho)
            if len(tentativas) == 1:
                raise _ComErro(-2147024809, "Parâmetro incorreto.", None, None)
            return object()

    class Mail:
        Attachments = Anexos()

    enviador = EnviadorOutlook.__new__(EnviadorOutlook)  # sem abrir o Outlook
    enviador._anexar(Mail(), pdf, "a apresentação")
    assert tentativas[0] == str(pdf)
    assert Path(tentativas[1]).parent == tmp_path / "tmp" / "disparo_laminas"
    assert Path(tentativas[1]).read_bytes() == b"%PDF"


def test_disparo_para_apos_erros_seguidos(hist):
    contatos = [Contato(i, "A", f"N{i}", f"n{i}@a.com", enviar=True) for i in range(6)]
    enviador = _Enviador(falhar={c.email for c in contatos})
    d, fila = _disparo(contatos, hist, enviador, _Relogio())
    d.run()
    erros = [r for r in hist.listar() if r["status"] == ERRO]
    assert len(erros) == 3
    assert "3 erros seguidos" in _eventos(fila)[-1][1]["motivo"]


class _MailFalso:
    """E-mail do Outlook falso: o Send só funciona depois de abrir a janela."""

    def __init__(self, exigir_janela=True):
        self.Attachments = self
        self.anexos, self.aberto, self.enviado = [], False, False
        self.exigir_janela = exigir_janela
        self.HTMLBody = "<html><body><p>Assinatura</p></body></html>"
        self.GetInspector = None

    def Add(self, caminho):  # Attachments.Add
        self.anexos.append(caminho)

        class _PA:
            def SetProperty(self, *a):
                pass

        class _Anexo:
            PropertyAccessor = _PA()
        return _Anexo()

    def Display(self, modal=False):
        self.aberto = True

    def Send(self):
        if self.exigir_janela and not self.aberto:
            raise _ComErro(-2147024809, "Parâmetro incorreto.", None, None)
        self.enviado = True


def _outlook_falso(exigir_janela):
    from laminas.envio import EnviadorOutlook
    criados = []

    class _App:
        def CreateItem(self, tipo):
            criados.append(_MailFalso(exigir_janela))
            return criados[-1]
    e = EnviadorOutlook.__new__(EnviadorOutlook)  # sem abrir o Outlook de verdade
    e.app, e.conta, e.incluir_assinatura, e.largura_max = _App(), None, True, 560
    return e, criados


def test_send_recusado_envia_abrindo_a_janela_como_a_macro(tmp_path):
    from laminas.mensagem import Email
    img = tmp_path / "l.png"
    img.write_bytes(PNG_1x1)
    email = Email(para="a@a.com", assunto="Oi", texto_inicio="Olá", texto_fim="", imagem=img)

    e, criados = _outlook_falso(exigir_janela=True)
    obs = e.enviar(email)
    assert criados[0].enviado and criados[0].aberto and "abrindo a janela" in obs

    e, criados = _outlook_falso(exigir_janela=False)
    assert e.enviar(email) == "" and not criados[0].aberto  # Send direto: sem janela


def test_observacao_do_envio_vai_para_o_historico(hist):
    class _EnviadorComObs(_Enviador):
        def enviar(self, email):
            super().enviar(email)
            return "Enviado abrindo a janela do e-mail (como a macro)"
    contatos = [Contato(2, "A", "Ana", "ana@a.com", enviar=True)]
    d, fila = _disparo(contatos, hist, _EnviadorComObs(), _Relogio())
    d.run()
    r = hist.listar()[0]
    assert r["status"] == ENVIADO and "como a macro" in r["erro"]


# -- versão ------------------------------------------------------------------------

def test_versao_documentada_e_gravada_no_historico(hist):
    import re
    from laminas import __version__
    historico_versoes = (Path(__file__).resolve().parent.parent / "disparo_laminas"
                         / "HISTORICO_VERSOES.md").read_text(encoding="utf-8")
    primeira = re.search(r"^## (\d+\.\d+\.\d+)", historico_versoes, flags=re.M).group(1)
    assert primeira == __version__, "suba __version__ e descreva a entrega em HISTORICO_VERSOES.md"
    hist.registrar(email="a@a.com")
    assert hist.listar()[0]["versao"] == __version__


# -- dados fora da pasta do programa ------------------------------------------------

def test_dados_ficam_fora_da_pasta_do_programa_e_migram(tmp_path, monkeypatch):
    import importlib
    monkeypatch.setenv("DISPARO_LAMINAS_DADOS", str(tmp_path / "appdata"))
    cfg = importlib.reload(cfgmod)
    try:
        assert cfg.PASTA_DADOS == tmp_path / "appdata"
        assert cfg.ARQUIVO_HISTORICO.parent == cfg.ARQUIVO_CONFIG.parent == cfg.PASTA_DADOS
        assert cfg.PASTA_PROGRAMA not in cfg.PASTA_DADOS.parents

        programa = tmp_path / "programa"
        (programa / "dados").mkdir(parents=True)
        (programa / "config.json").write_text('{"limite_diario": 50}', encoding="utf-8")
        Historico(programa / "dados" / "historico_envios.db").registrar(email="a@a.com")
        assert sorted(cfg.migrar_dados_antigos(programa, cfg.PASTA_DADOS)) == \
            ["config.json", "historico_envios.db"]
        assert cfg.carregar(cfg.ARQUIVO_CONFIG)["limite_diario"] == 50
        assert Historico(cfg.ARQUIVO_HISTORICO).data_envio("a@a.com") is not None
        # não sobrescreve o que já está na pasta de dados
        (programa / "config.json").write_text('{"limite_diario": 10}', encoding="utf-8")
        assert cfg.migrar_dados_antigos(programa, cfg.PASTA_DADOS) == []
        assert cfg.carregar(cfg.ARQUIVO_CONFIG)["limite_diario"] == 50
    finally:
        monkeypatch.delenv("DISPARO_LAMINAS_DADOS")
        importlib.reload(cfgmod)


# -- 1.7.0: campanhas salvas, remetente/subdomínio, backup ------------------------

def test_campanhas_salvas_salvar_listar_comparar_excluir(tmp_path):
    from laminas import salvas
    cfg = cfgmod.carregar(tmp_path / "c.json")
    lamina = tmp_path / "CORPORATIVO" / "prazo70.jpeg"
    lamina.parent.mkdir()
    lamina.write_bytes(PNG_1x1)
    estado = {"pasta": "CORPORATIVO", "laminas": [lamina], "alternar": False,
              "anexo": "Institucional.pdf",
              "mensagem": {"assunto": "Oi {nome}", "texto_inicio": "a", "texto_fim": "b"}}
    nome = salvas.salvar(cfg, "  Corporativo   Prazo 70  ", estado)
    assert nome == "Corporativo Prazo 70" and salvas.listar(cfg) == [nome]
    assert salvas.existente(cfg, "corporativo prazo 70") == nome
    assert salvas.mesmo_conteudo(salvas.obter(cfg, nome), estado)
    assert not salvas.mesmo_conteudo(salvas.obter(cfg, nome), {**estado, "anexo": ""})
    # mesmo nome com outra caixa substitui em vez de duplicar
    salvas.salvar(cfg, "CORPORATIVO PRAZO 70", {**estado, "alternar": True})
    assert salvas.listar(cfg) == [nome] and salvas.obter(cfg, nome)["alternar"] is True
    # gravada no config.json e lida de volta
    cfgmod.salvar(cfg, tmp_path / "c.json")
    assert salvas.listar(cfgmod.carregar(tmp_path / "c.json")) == [nome]
    for invalido in ("", "   ", "(nova campanha)", "x" * 61):
        with pytest.raises(salvas.NomeInvalido):
            salvas.normalizar_nome(invalido)
    salvas.excluir(cfg, nome)
    assert salvas.listar(cfg) == []


def test_campanha_salva_acha_lamina_que_mudou_de_pasta(tmp_path):
    from laminas import salvas
    nova = tmp_path / "novo_lugar" / "prazo70.jpeg"
    nova.parent.mkdir()
    nova.write_bytes(PNG_1x1)
    salva = {"laminas": [str(tmp_path / "antigo" / "prazo70.jpeg"), str(tmp_path / "sumiu.jpeg")]}
    encontradas, faltando = salvas.localizar_laminas(salva, [nova])
    assert encontradas == [nova] and faltando == ["sumiu.jpeg"]


def test_senha_smtp_nao_fica_em_texto_puro():
    from laminas import segredo
    protegida = segredo.proteger("minha-senha")
    assert "minha-senha" not in protegida and segredo.revelar(protegida) == "minha-senha"
    assert segredo.proteger("") == "" and segredo.revelar("legado") == "legado"


def test_config_antiga_com_conta_outlook_vira_modo_conta(tmp_path):
    caminho = tmp_path / "c.json"
    caminho.write_text('{"conta_outlook": "comercial@prospeccao.x.com.br"}', encoding="utf-8")
    cfg = cfgmod.carregar(caminho)
    assert cfg["modo_envio"] == "outlook_conta"
    assert cfg["remetente_email"] == "comercial@prospeccao.x.com.br"
    assert cfgmod.descricao_remetente(cfg) == "comercial@prospeccao.x.com.br (Outlook, conta)"
    assert cfgmod.carregar(tmp_path / "nao-existe.json")["modo_envio"] == "outlook"


def test_avisos_de_subdominio():
    base = dict(cfgmod.DEFAULTS)
    assert cfgmod.normalizar_subdominio(" https://Prospeccao.X.com.br/ ") == "prospeccao.x.com.br"
    ok = {**base, "modo_envio": "outlook_conta", "subdominio_envio": "prospeccao.x.com.br",
          "remetente_email": "ana@prospeccao.x.com.br"}
    assert cfgmod.avisos_remetente(ok) == []
    fora = {**ok, "remetente_email": "ana@x.com.br"}
    assert any("não é do subdomínio" in a for a in cfgmod.avisos_remetente(fora))
    padrao = {**ok, "modo_envio": "outlook"}
    assert any("NÃO sairão pelo subdomínio" in a for a in cfgmod.avisos_remetente(padrao))
    smtp = {**ok, "modo_envio": "smtp"}
    assert any("SMTP sem servidor, usuário, senha" in a for a in cfgmod.avisos_remetente(smtp))


class _SMTPFalso:
    """Servidor SMTP falso: guarda as mensagens enviadas."""
    enviados: list = []
    ultimo = None

    def __init__(self, host, porta, timeout=None, context=None):
        self.host, self.porta, self.tls, self.logado = host, porta, False, None
        _SMTPFalso.ultimo = self

    def ehlo(self):
        pass

    def starttls(self, context=None):
        self.tls = True

    def login(self, usuario, senha):
        if senha != "certa":
            import smtplib
            raise smtplib.SMTPAuthenticationError(535, b"auth failed")
        self.logado = usuario

    def send_message(self, msg, from_addr=None, to_addrs=None):
        _SMTPFalso.enviados.append((msg, from_addr, list(to_addrs)))

    def noop(self):
        return (250, b"ok")

    def quit(self):
        pass

    def close(self):
        pass


def test_envio_smtp_pelo_subdominio(tmp_path, monkeypatch):
    import smtplib
    from laminas.envio import EnviadorSMTP, ErroEnvio, criar_enviador
    from laminas import segredo
    monkeypatch.setattr(smtplib, "SMTP", _SMTPFalso)
    _SMTPFalso.enviados = []
    img = tmp_path / "lamina.png"
    img.write_bytes(PNG_1x1)
    pdf = tmp_path / "Institucional.pdf"
    pdf.write_bytes(b"%PDF-1.4")
    cfg = {**cfgmod.DEFAULTS, "modo_envio": "smtp", "smtp_host": "smtp.prospeccao.x.com.br",
           "smtp_porta": 587, "smtp_seguranca": "starttls", "smtp_usuario": "comercial",
           "smtp_senha": segredo.proteger("certa"), "remetente_email": "comercial@prospeccao.x.com.br",
           "remetente_nome": "Filipe | Tivolitur", "responder_para": "filipe@x.com.br",
           "assinatura_smtp": "Filipe Veronezzi\nTivolitur"}
    enviador = criar_enviador(cfg, tmp_path / "sim")
    assert isinstance(enviador, EnviadorSMTP)
    c = Contato(2, "Acme", "Ana", "ana@acme.com", cc=True)
    email = montar(c, {"assunto": "Oi {nome}", "texto_inicio": "Olá {nome}", "texto_fim": "Até"},
                   imagem=img, cc="ricardo@x.com", cco="arquivo@x.com", anexos=[pdf])
    assert enviador.enviar(email) == ""
    msg, de, para = _SMTPFalso.enviados[0]
    assert _SMTPFalso.ultimo.tls and _SMTPFalso.ultimo.logado == "comercial"
    assert de == "comercial@prospeccao.x.com.br"
    assert para == ["ana@acme.com", "ricardo@x.com", "arquivo@x.com"]  # CCO só no envelope
    assert "Bcc" not in msg and msg["Cc"] == "ricardo@x.com" and msg["Reply-To"] == "filipe@x.com.br"
    assert msg["From"] == "Filipe | Tivolitur <comercial@prospeccao.x.com.br>"
    assert msg["Message-ID"].endswith("@prospeccao.x.com.br>")
    html = msg.get_body(preferencelist=("html",)).get_content()
    assert "Olá Ana" in html and "cid:" in html and "Filipe Veronezzi" in html
    inline = [p for p in msg.walk() if p.get("Content-ID")]
    assert inline and f"cid:{inline[0]['Content-ID'][1:-1]}" in html
    anexos = [p.get_filename() for p in msg.iter_attachments()]
    assert "Institucional.pdf" in anexos
    # senha errada -> mensagem clara
    cfg["smtp_senha"] = segredo.proteger("errada")
    with pytest.raises(ErroEnvio, match="recusou o usuário/senha"):
        criar_enviador(cfg, tmp_path / "sim").enviar(email)
    with pytest.raises(ErroEnvio, match="sem servidor"):
        criar_enviador({**cfg, "smtp_host": ""}, tmp_path / "sim")


def test_outlook_enviar_como_e_responder_para():
    from laminas.envio import EnviadorOutlook
    e, criados = _outlook_falso(exigir_janela=False)
    e.enviar_como, e.responder_para = "comercial@prospeccao.x.com.br", "filipe@x.com.br"

    class _Respostas:
        def __init__(self):
            self.lista = []

        def Add(self, endereco):
            self.lista.append(endereco)
    _MailFalso.ReplyRecipients = _Respostas()
    try:
        from laminas.mensagem import Email
        e.enviar(Email(para="a@a.com", assunto="Oi", texto_inicio="Olá", texto_fim=""))
        assert criados[0].SentOnBehalfOfName == "comercial@prospeccao.x.com.br"
        assert _MailFalso.ReplyRecipients.lista == ["filipe@x.com.br"]
    finally:
        del _MailFalso.ReplyRecipients


def test_backup_diario_e_por_versao_preserva_historico(tmp_path):
    from datetime import datetime as dt
    dados = tmp_path / "dados"
    h = Historico(dados / "historico_envios.db")
    h.registrar(email="ja@enviado.com")
    cfgmod.salvar({"limite_diario": 50}, dados / "config.json")
    b1 = cfgmod.backup_dados("1.7.0", dados, agora=dt(2026, 10, 30, 9))
    assert b1 and Historico(b1 / "historico_envios.db").data_envio("ja@enviado.com") is not None
    assert (b1 / "config.json").exists()
    assert cfgmod.backup_dados("1.7.0", dados, agora=dt(2026, 10, 30, 15)) is None  # mesmo dia/versão
    assert cfgmod.backup_dados("1.7.1", dados, agora=dt(2026, 10, 30, 16)) is not None  # versão nova
    from datetime import timedelta
    for dia in range(40):  # um backup por dia: mantém só os mais recentes
        cfgmod.backup_dados("1.7.1", dados, agora=dt(2026, 11, 1, 9) + timedelta(days=dia))
    pastas = sorted(p.name for p in (dados / "backups").iterdir() if p.is_dir())
    assert len(pastas) == cfgmod.MAX_BACKUPS and pastas[-1].startswith("20261210")


def test_historico_registra_remetente(hist):
    hist.registrar(email="a@a.com", remetente="comercial@prospeccao.x.com.br (SMTP)")
    assert hist.listar()[0]["remetente"] == "comercial@prospeccao.x.com.br (SMTP)"


# -- 1.7.1: Outlook na thread da tela, reconexão e modo janela -------------------

def test_disparo_envia_pela_thread_da_tela(hist):
    import threading
    import time as _time
    from laminas.executor import EnviadorNaTela, ExecutorPrincipal
    executor = ExecutorPrincipal()
    principal = threading.current_thread()
    threads_usadas = []

    class _EnviadorDaTela:
        def __init__(self):
            threads_usadas.append(threading.current_thread())

        def enviar(self, email):
            threads_usadas.append(threading.current_thread())
            if email.para == "falha@a.com":
                raise RuntimeError("caixa cheia")
            return ""

    contatos = [Contato(2, "A", "Ana", "ana@a.com", enviar=True),
                Contato(3, "B", "Bia", "falha@a.com", enviar=True),
                Contato(4, "C", "Caio", "caio@c.com", enviar=True)]
    d, fila = _disparo(contatos, hist, None, _Relogio(),
                       criar_enviador=lambda: EnviadorNaTela(executor, _EnviadorDaTela))
    d.start()  # o disparo em segundo plano; a "tela" é esta thread
    limite = _time.monotonic() + 10
    while d.is_alive() and _time.monotonic() < limite:
        executor.processar_pendentes()
        _time.sleep(0.005)
    assert not d.is_alive()
    assert threads_usadas and all(t is principal for t in threads_usadas)
    status = {r["email"]: (r["status"], r["erro"]) for r in hist.listar()}
    assert status["ana@a.com"][0] == ENVIADO and status["caio@c.com"][0] == ENVIADO
    assert status["falha@a.com"] == (ERRO, "caixa cheia")


def test_outlook_reconecta_quando_a_ligacao_cai():
    from laminas.mensagem import Email
    e, criados = _outlook_falso(exigir_janela=False)
    app_bom = e.app

    class _AppCaido:
        def CreateItem(self, tipo):
            raise _ComErro(-2147417848, "O objeto chamado foi desconectado de seus clientes.", None, None)
    e.app = _AppCaido()

    def _conectar():
        e.app, e.conta = app_bom, None
    e._conectar = _conectar
    obs = e.enviar(Email(para="a@a.com", assunto="Oi", texto_inicio="Olá", texto_fim=""))
    assert criados[0].enviado and "Outlook reconectado" in obs and "0x80010108" in obs


def test_depois_de_um_send_recusado_os_proximos_ja_abrem_a_janela():
    from laminas.mensagem import Email
    e, criados = _outlook_falso(exigir_janela=True)
    email = Email(para="a@a.com", assunto="Oi", texto_inicio="Olá", texto_fim="")
    tentativas_diretas = []
    _send = _MailFalso.Send

    def _send_contando(self):
        if not self.aberto:
            tentativas_diretas.append(self)
        return _send(self)
    _MailFalso.Send = _send_contando
    try:
        e.enviar(email)
        e.enviar(email)
        e.enviar(email)
    finally:
        _MailFalso.Send = _send
    assert all(m.enviado for m in criados) and e.usar_janela
    assert len(tentativas_diretas) == 1  # só o 1º tentou o envio direto
