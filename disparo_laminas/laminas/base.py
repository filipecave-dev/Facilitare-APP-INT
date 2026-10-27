"""Leitura da lista de envio e verificação de duplicados.

A base é uma planilha (``.xlsx``/``.xlsm``/``.csv``) com, no mínimo, as colunas
**Empresa**, **Nome** e **E-mail** (os títulos são reconhecidos por
aproximação: "Nome Completo", "E-mail Principal" etc.). Uma coluna opcional
**Cópia**/**CC** com "S"/"X"/"Sim" já deixa o contato marcado para cópia.

Cada linha vira um :class:`Contato` com uma situação:

- ``pendente``   — pronto para envio;
- ``duplicado``  — o mesmo e-mail aparece numa linha anterior da base;
- ``ja_enviado`` — o e-mail já está no histórico de enviados;
- ``invalido``   — e-mail vazio/mal formado ou faltando nome/empresa;
- ``enviado``    — enviado nesta sessão.
"""

from __future__ import annotations

import csv
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import Iterable, Optional

PENDENTE = "pendente"
DUPLICADO = "duplicado"
JA_ENVIADO = "ja_enviado"
INVALIDO = "invalido"
ENVIADO = "enviado"

_RE_EMAIL = re.compile(r"^[A-Za-z0-9._%+'\-]+@[A-Za-z0-9\-]+(\.[A-Za-z0-9\-]+)*\.[A-Za-z]{2,}$")
_SIM = {"s", "sim", "x", "1", "true", "verdadeiro", "cc", "copia", "yes"}


@dataclass
class Contato:
    linha: int
    empresa: str
    nome: str
    email: str
    situacao: str = PENDENTE
    detalhe: str = ""
    enviar: bool = False
    cc: bool = False
    # Reenvio autorizado manualmente para um e-mail que já recebeu lâmina.
    forcar: bool = False

    @property
    def email_norm(self) -> str:
        return normalizar_email(self.email)

    @property
    def selecionavel(self) -> bool:
        return self.situacao in (PENDENTE, JA_ENVIADO)


@dataclass
class Resumo:
    total: int = 0
    pendentes: int = 0
    ja_enviados: int = 0
    invalidos: int = 0
    # e-mail -> números das linhas em que aparece (só os repetidos)
    duplicados: dict[str, list[int]] = field(default_factory=dict)


# -- limpeza ---------------------------------------------------------------

def _sem_acento(texto: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", texto)
                   if not unicodedata.combining(c))


def limpar_texto(valor) -> str:
    """Remove lixo de exportação (``_x000D_``, quebras, tabs) e espaços extras."""
    if valor is None:
        return ""
    texto = str(valor).replace("_x000D_", " ")
    return " ".join(texto.split())


def limpar_email(valor) -> str:
    texto = limpar_texto(valor)
    if texto.lower().startswith("mailto:"):
        texto = texto[7:]
    return texto.strip(" <>;,").replace(" ", "")


def normalizar_email(valor) -> str:
    return limpar_email(valor).lower()


def email_valido(email: str) -> bool:
    return bool(_RE_EMAIL.match(email or ""))


def _titulo(valor) -> str:
    return _sem_acento(limpar_texto(valor)).lower()


# -- leitura de planilhas --------------------------------------------------

def listar_abas(caminho: str) -> list[str]:
    if Path(caminho).suffix.lower() == ".csv":
        return []
    from openpyxl import load_workbook
    wb = load_workbook(caminho, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _linhas_csv(caminho: str) -> list[tuple]:
    bruto = Path(caminho).read_bytes()
    for codificacao in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            texto = bruto.decode(codificacao)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialeto = csv.Sniffer().sniff(texto[:4096], delimiters=";,\t")
    except csv.Error:
        dialeto = csv.excel
        dialeto.delimiter = ";"
    return [tuple(r) for r in csv.reader(texto.splitlines(), dialeto)]


def ler_linhas(caminho: str, aba: Optional[str] = None) -> list[tuple]:
    """Todas as linhas (tuplas de valores) da aba ou do CSV."""
    if Path(caminho).suffix.lower() == ".csv":
        return _linhas_csv(caminho)
    from openpyxl import load_workbook
    wb = load_workbook(caminho, read_only=True, data_only=True)
    try:
        ws = wb[aba] if aba and aba in wb.sheetnames else wb.worksheets[0]
        return [tuple(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def mapear_colunas(linhas: list[tuple]) -> tuple[int, dict[str, int]]:
    """Acha a linha de títulos (entre as 10 primeiras) e o índice de cada
    coluna conhecida. Retorna ``(indice_da_linha_de_titulo, colunas)``."""
    for i, linha in enumerate(linhas[:10]):
        titulos = [_titulo(v) for v in linha]
        colunas: dict[str, int] = {}
        for j, t in enumerate(titulos):
            if not t:
                continue
            eh_copia = "copia" in t or t in ("cc", "cco", "bcc") or "(cc)" in t
            if eh_copia:
                if "oculta" not in t and t not in ("cco", "bcc"):
                    colunas.setdefault("copia", j)
            elif "mail" in t:
                colunas.setdefault("email", j)
            elif "empresa" in t:
                colunas.setdefault("empresa", j)
            elif "nome" in t:
                colunas.setdefault("nome", j)
            elif t.startswith("data"):
                colunas.setdefault("data", j)
            elif t.startswith("hora"):
                colunas.setdefault("hora", j)
            elif "status" in t or "situacao" in t:
                colunas.setdefault("status", j)
        if "email" in colunas:
            return i, colunas
    raise ValueError(
        "Não encontrei a coluna de E-mail nas primeiras linhas da planilha. "
        "Confira se a aba tem os títulos Empresa / Nome / E-mail."
    )


def _valor(linha: tuple, colunas: dict[str, int], chave: str):
    j = colunas.get(chave)
    return linha[j] if j is not None and j < len(linha) else None


def registros(caminho: str, aba: Optional[str] = None) -> list[dict]:
    """Linhas da planilha como dicionários (empresa, nome, email, copia,
    data, status, linha). Linhas totalmente vazias são ignoradas."""
    linhas = ler_linhas(caminho, aba)
    if not linhas:
        return []
    inicio, colunas = mapear_colunas(linhas)
    saida = []
    for n, linha in enumerate(linhas[inicio + 1:], start=inicio + 2):
        reg = {
            "linha": n,
            "empresa": limpar_texto(_valor(linha, colunas, "empresa")),
            "nome": limpar_texto(_valor(linha, colunas, "nome")),
            "email": limpar_email(_valor(linha, colunas, "email")),
            "copia": limpar_texto(_valor(linha, colunas, "copia")),
            "data": _data_hora(_valor(linha, colunas, "data"), _valor(linha, colunas, "hora")),
            "status": limpar_texto(_valor(linha, colunas, "status")),
        }
        if reg["email"] or reg["nome"] or reg["empresa"]:
            saida.append(reg)
    return saida


def _data_hora(data, hora) -> Optional[datetime]:
    if isinstance(data, datetime):
        base = data
    elif isinstance(data, date):
        base = datetime.combine(data, time())
    elif isinstance(data, str) and data.strip():
        base = None
        for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S"):
            try:
                base = datetime.strptime(data.strip(), formato)
                break
            except ValueError:
                continue
        if base is None:
            return None
    else:
        return None
    if isinstance(hora, time) and base.time() == time():
        base = datetime.combine(base.date(), hora)
    elif isinstance(hora, str) and base.time() == time():
        try:
            base = datetime.combine(base.date(), time.fromisoformat(hora.strip()))
        except ValueError:
            pass
    return base


# -- análise ---------------------------------------------------------------

def analisar(regs: Iterable[dict], ja_enviados: dict[str, Optional[str]]) -> tuple[list[Contato], Resumo]:
    """Classifica cada linha e aponta duplicados.

    ``ja_enviados`` mapeia e-mail normalizado -> data/hora do envio anterior
    (texto ``AAAA-MM-DD HH:MM:SS`` ou ``None`` quando a data é desconhecida).
    """
    contatos: list[Contato] = []
    resumo = Resumo()
    primeira_linha: dict[str, int] = {}
    for reg in regs:
        c = Contato(linha=reg["linha"], empresa=reg["empresa"], nome=reg["nome"],
                    email=limpar_email(reg["email"]), cc=_titulo(reg.get("copia")) in _SIM)
        norm = c.email_norm
        if not c.email:
            c.situacao, c.detalhe = INVALIDO, "Sem e-mail"
        elif not email_valido(c.email):
            c.situacao, c.detalhe = INVALIDO, "E-mail inválido"
        elif norm in primeira_linha:
            c.situacao = DUPLICADO
            c.detalhe = f"Duplicado da linha {primeira_linha[norm]}"
            resumo.duplicados.setdefault(norm, [primeira_linha[norm]]).append(c.linha)
        else:
            primeira_linha[norm] = c.linha
            if norm in ja_enviados:
                c.situacao = JA_ENVIADO
                c.detalhe = "Já enviado" + _quando(ja_enviados[norm])
            elif not c.nome or not c.empresa:
                c.situacao = INVALIDO
                falta = " e ".join(x for x, v in (("nome", c.nome), ("empresa", c.empresa)) if not v)
                c.detalhe = f"Falta {falta}"
            else:
                c.situacao, c.enviar = PENDENTE, True
        if not c.selecionavel or c.situacao == JA_ENVIADO:
            c.cc = False
        contatos.append(c)

    resumo.total = len(contatos)
    resumo.pendentes = sum(c.situacao == PENDENTE for c in contatos)
    resumo.ja_enviados = sum(c.situacao == JA_ENVIADO for c in contatos)
    resumo.invalidos = sum(c.situacao == INVALIDO for c in contatos)
    return contatos, resumo


def _quando(data_hora: Optional[str]) -> str:
    if not data_hora:
        return " (envio anterior)"
    try:
        return " em " + datetime.fromisoformat(data_hora).strftime("%d/%m/%Y %H:%M")
    except ValueError:
        return f" em {data_hora}"
